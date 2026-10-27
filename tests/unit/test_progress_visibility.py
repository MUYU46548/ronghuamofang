# -*- coding: utf-8 -*-
"""可观测性 / 产物可见性（2026-10-26，用户 13 条反馈的确定性修复）。

## 为什么需要它

一次真实试跑之后，用户列了 13 条体验缺陷。其中五条**不是设计取舍，是 bug**，
而且都有同一个特征：界面显示的东西与盘上的事实不一致，且全程无告警。

| 现象（用户原话） | 真因 |
|---|---|
| 「我受不了那个永恒的 29%」 | `stage_status` 从不返回 `chapters_done`，前端拿 undefined 算进度 → running 阶段贡献恒 0 |
| 「阶段 8 在客户端内直接不存在」 | `build_state` 写死 `range(1, 8)`，阶段 8 被整个排除 |
| 「output (1 篇 · 0 字)」 | `summarize_artifacts` 遇到 .docx 直接 continue → chars 恒 0 |
| 「第10章潜七十二，中间没有空格」 | 章标题无归一化，`TITLE_RE` 对「章」后内容全宽容 |
| 「每章只有两千多字，作为小说严重不合格」 | 安装态 config 只补缺不覆盖 → target_words 停在旧默认 [2000,5500] |

## 判据

A. `stage_status` 返回章级进度（进度键 + 盘上产物文件数取大者），且 `started_at` 可透出。
B. `build_state().stages` 覆盖 1-8（含 Markdown 导出），不再有"看不见的阶段"。
C. `summarize_artifacts` 对 .docx **真数正文**，不再谎报 0 字。
D. 章标题归一化：粘连 → 加空格；已规范 → 幂等；无章名 → 不留尾随空格；正文「第N章」不动。
E. `chapter.title_format` 可读可写，缺 `{n}` 一律 400。

全程离线、零 LLM、零费用。
用法：python tests/unit/test_progress_visibility.py
"""
import io
import json
import os
import sys
import tempfile
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
os.environ.pop("NF_ROOT", None)

PASS, FAIL = [], []


def check(name, cond, detail=""):
    (PASS if cond else FAIL).append(name)
    print("  [%s] %s%s" % ("PASS" if cond else "FAIL", name,
                           ("  → " + str(detail)[:240]) if detail else ""))


def _write_progress(root, stages):
    p = Path(root) / "data" / "state" / "progress.json"
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps({"project": "t", "version": 3, "stages": stages,
                             "budget": {}}, ensure_ascii=False), encoding="utf-8")


def _seed_workspace(root):
    """最小工作区：config/project.yaml + 各阶段产物目录。"""
    (Path(root) / "config").mkdir(parents=True, exist_ok=True)
    (Path(root) / "config" / "project.yaml").write_text(
        "book:\n  name: 探针书\n  chapters: 10\n", encoding="utf-8")
    (Path(root) / "config" / "system.yaml").write_text(
        "chapter:\n  target_words: [5000, 9000]\n", encoding="utf-8")
    for rel in ("data/chapters/raw", "data/chapters/checked",
                "data/chapters/refined", "data/outline/chapters"):
        d = Path(root) / rel
        d.mkdir(parents=True, exist_ok=True)
        for i in range(1, 4):
            (d / ("%02d.md" % i)).write_text("## 第%d章 测试\n\n正文。\n" % i,
                                             encoding="utf-8")


def main():
    import nf_api

    print("【A】阶段内章级进度（用户「永恒 29%」）")
    tmp = tempfile.mkdtemp(prefix="nf_prog_")
    try:
        _seed_workspace(tmp)
        _write_progress(tmp, {"4": {"status": "running", "completed_chapters": [1, 2],
                                    "started_at": "2026-10-26T10:00:00",
                                    "current_chapter": 3, "current_chars": 1200},
                              "5": {"status": "running"}})
        nf_api._set_root(Path(tmp))
        st4 = nf_api.stage_status(nf_api.ProgressManager(Path(tmp) / "data/state/progress.json"), 4, 10)
        check("阶段 4 返回 chapters_done", st4.get("chapters_done") == 3,
              st4)   # 进度键 2 章 vs 盘上 3 章 → 取大者（盘上事实优先）
        check("阶段 4 返回 total_chapters", st4.get("total_chapters") == 10, st4)
        check("阶段 4 返回 started_at（前端好算「已跑多久」）",
              st4.get("started_at") == "2026-10-26T10:00:00", st4)
        check("阶段 4 在写时透出 current_chapter（「正在写第 3 章」）",
              st4.get("current_chapter") == 3, st4)
        check("阶段 4 在写时透出 current_chars（「已写 1200 字」）",
              st4.get("current_chars") == 1200, st4)
        st5 = nf_api.stage_status(nf_api.ProgressManager(Path(tmp) / "data/state/progress.json"), 5, 10)
        check("阶段 5 无进度键也能从 checked/ 数出章数（不再恒 0）",
              st5.get("chapters_done") == 3, st5)
    finally:
        nf_api._set_root(ROOT)

    print("\n【B】阶段 8 必须出现在 /state（用户「客户端内直接不存在」）")
    tmp2 = tempfile.mkdtemp(prefix="nf_state_")
    try:
        _seed_workspace(tmp2)
        _write_progress(tmp2, {"8": {"status": "done"}})
        nf_api._set_root(Path(tmp2))
        state = nf_api.build_state()
        nums = [s["stage"] for s in state["stages"]]
        check("stages 覆盖 1-8", nums == list(range(1, 9)), nums)
        s8 = [s for s in state["stages"] if s["stage"] == 8][0]
        check("阶段 8 的状态可读", s8["status"] == "done", s8)
        check("阶段 8 也带章级字段（前端同一套渲染）",
              "chapters_done" in s8 and "total_chapters" in s8, s8)
        # 2026-10-27 用户实测：「阶段8无法放行（提示stage须为1-7，代码打架）」。
        # 阶段 8 出现在 /state 且带「确认」按钮，审批端点就必须收它。
        ok8, msg8 = nf_api.act_approve(8, False)
        check("阶段 8 可审批（不再被 1-7 挡在门外）", ok8 is True, msg8)
        pm8 = nf_api.ProgressManager(Path(tmp2) / "data/state/progress.json")
        check("阶段 8 审批位真的落到 progress", pm8.is_approved(8) is True,
              pm8.data["stages"]["8"])
    finally:
        nf_api._set_root(ROOT)

    print("\n【B2】前端阶段表 / 复制任务 / 成品入口（防回潮）")
    app_vue = (ROOT / "console" / "src" / "App.vue").read_text(encoding="utf-8")
    check("前端 STAGE_NAMES 收录阶段 8（Markdown 导出）",
          "Markdown 导出" in app_vue, "App.vue 里找不到阶段 8 的名字")
    check("阶段 8 在 STAGE_NAMES 的 8 号位",
          '8: "Markdown 导出"' in app_vue or "8: 'Markdown 导出'" in app_vue)
    check("复制任务按钮判据是 STAGE_TASK_INFO（阶段 7/8 有任务、无模板）",
          "agentMode && STAGE_TASK_INFO[s.stage]" in app_vue,
          "若退回 STAGE_MAIN_PROMPT[s.stage]，阶段 7 的空串又会让按钮消失")
    check("STAGE_TASK_INFO 收录阶段 7（Word 成品）", "7: { name:" in app_vue)
    check("STAGE_TASK_INFO 收录阶段 8（Markdown 导出）", "8: { name:" in app_vue)
    check("顶栏有「打开成品」入口（不必再靠命令面板找 docx）",
          "📕 成品" in app_vue)
    check("导出记录读盘上事实（/exports/history）", "/exports/history" in app_vue)
    # 2026-10-27：单章重写显示「本章 %」+ 段群明细，禁止退回「整体 99% 不动」。
    check("前端按「本章内进度」计算（stageFraction）", "function stageFraction" in app_vue)
    check("前端显示段群进度（第 g/G 群 · 本群 X/Y 字）",
          "本群 " in app_vue and "current_group_quota" in app_vue)
    check("单章重写显示「本章 %」而不是全书百分比",
          "scopeChapterPct" in app_vue and "本章 " in app_vue)
    check("重写完成提示带字节数与群字数（/chapters/content 的 bytes/parts）",
          "d.bytes" in app_vue and "d.parts" in app_vue)

    print("\n【C】docx 字数不再谎报 0（用户「1 篇 · 0 字」）")
    from utils.stage_artifacts import summarize_artifacts
    tmp3 = tempfile.mkdtemp(prefix="nf_docx_")
    try:
        out = Path(tmp3) / "output"
        out.mkdir(parents=True)
        try:
            import docx as _docx
            d = _docx.Document()
            d.add_heading("第一章", level=1)
            d.add_paragraph("这是一段有正文的段落，用来验证字数统计不是 0。")
            d.save(str(out / "书_完整版.docx"))
            have_docx = True
        except Exception as e:                               # noqa: BLE001
            have_docx = False
            check(True, "跳过 docx 断言（环境无 python-docx）: " + str(e)[:80])
        if have_docx:
            items = summarize_artifacts(7, tmp3)
            it = items[0]
            check("docx 目录 files 计数正常", it["files"] == 1, it)
            check("docx 正文字数 > 0（不再显示 0 字）", it["chars"] > 10, it)
            check("字数口径含标题与段落", it["chars"] >= len("第一章") + 10, it)
    finally:
        pass

    print("\n【D】章标题格式归一化（用户「第10章潜七十二」）")
    from utils.text_norm import normalize_chapter_title
    cases = [
        ("## 第10章潜七十二\n\n正文。", "## 第10章 潜七十二", 1),
        ("## 第1章 归港\n\n正文。", "## 第1章 归港", 0),          # 幂等
        ("## 第3章\n\n正文。", "## 第3章", 0),                    # 无章名不留尾空格
        ("### 第8章潜零前十二刻\n正文。", "### 第8章 潜零前十二刻", 1),
        ("##第2章清单末行\n正文。", "## 第2章 清单末行", 1),      # 连 # 后空格都没有
    ]
    for src, want, want_n in cases:
        got, n = normalize_chapter_title(src)
        check("归一化 %r → %r" % (src.splitlines()[0], want),
              got.splitlines()[0] == want and n == want_n, got.splitlines()[0])
    body, n = normalize_chapter_title("这是正文，提到第10章潜七十二的说法。\n")
    check("正文里的「第N章」不动（只改标题行）",
          "第10章潜七十二" in body and n == 0, body)
    again, n2 = normalize_chapter_title("## 第10章 潜七十二\n\n正文。")
    check("已是目标格式 → 幂等（不重复改写）", n2 == 0, again.splitlines()[0])

    print("\n【E】chapter.title_format 可读可写（用户要求「格式自定义」）")
    tmp4 = tempfile.mkdtemp(prefix="nf_tfmt_")
    try:
        _seed_workspace(tmp4)
        nf_api._set_root(Path(tmp4))
        from nf_api_domains import project as dom_project
        code, payload = dom_project.handle_config_title_format(None)
        check("读口 200", code == 200 and payload.get("ok"), payload)
        check("缺省回落 `第{n}章 {title}`",
              payload.get("current") == "第{n}章 {title}", payload.get("current"))
        code, payload = dom_project.handle_config_title_format_set(
            None, {"format": "第{n}章·{title}"})
        check("写口接受合法格式", code == 200 and payload.get("ok"), payload)
        code, payload = dom_project.handle_config_title_format_set(None, {"format": "无章号格式"})
        check("缺 {n} → 400 且说清原因", code == 400 and "{n}" in payload.get("error", ""),
              payload)
        code, payload = dom_project.handle_config_title_format_set(None, {"format": ""})
        check("空格式 → 400", code == 400, payload)
        # 读完应能拿到刚写进去的值（定向改写 + 回读）
        code, payload = dom_project.handle_config_title_format(None)
        check("写入后读回一致", payload.get("current") == "第{n}章·{title}",
              payload.get("current"))
    finally:
        nf_api._set_root(ROOT)

    print("\n【F】单章重写的进度分母 = 这一章（用户「已写1字，但进度99%」）")
    tmp5 = tempfile.mkdtemp(prefix="nf_rewrite_prog_")
    try:
        _seed_workspace(tmp5)
        # 重写第 5 章、加强模式第 2/5 群、本群 800/1400 字
        _write_progress(tmp5, {"4": {
            "status": "running", "rewrite_chapter": 5,
            "current_chapter": 5, "current_chars": 800,
            "current_group": 2, "total_groups": 5,
            "current_group_quota": 1400,
            "current_chars_at": "2026-10-27T15:00:00",
            "started_at": "2026-10-27T14:59:00"}})
        nf_api._set_root(Path(tmp5))
        pm = nf_api.ProgressManager(Path(tmp5) / "data/state/progress.json")
        st = nf_api.stage_status(pm, 4, 10)
        check("重写时分母是 1（不是全书 10 章）", st.get("total_chapters") == 1, st)
        check("重写时 chapters_done 只看这一章（未落盘 = 0）",
              st.get("chapters_done") == 0, st)
        check("透出 rewrite_chapter（前端可显示「重写第 5 章」）",
              st.get("rewrite_chapter") == 5, st)
        check("透出段群进度（第 2/5 群）",
              st.get("current_group") == 2 and st.get("total_groups") == 5, st)
        check("透出本群配额（1400）与最后更新时间",
              st.get("current_group_quota") == 1400
              and st.get("current_chars_at") == "2026-10-27T15:00:00", st)
        state5 = nf_api.build_state()
        check("build_state 透出 chapter_target_words（前端算本章 %）",
              state5.get("chapter_target_words") == [5000, 9000],
              state5.get("chapter_target_words"))

        # 非重写（正常全量）：分母回到全书章数，避免"修重写把正常进度改坏"
        _write_progress(tmp5, {"4": {
            "status": "running", "completed_chapters": list(range(1, 10)),
            "current_chapter": 5, "current_chars": 800}})
        st2 = nf_api.stage_status(
            nf_api.ProgressManager(Path(tmp5) / "data/state/progress.json"), 4, 10)
        check("非重写时仍按全书章数算（10）", st2.get("total_chapters") == 10, st2)
        check("非重写时不透出 rewrite_chapter", "rewrite_chapter" not in st2, st2)
    finally:
        nf_api._set_root(ROOT)

    print("\n" + "=" * 66)
    print("通过 %d / 失败 %d" % (len(PASS), len(FAIL)))
    if FAIL:
        for f in FAIL:
            print("  FAIL:", f)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
