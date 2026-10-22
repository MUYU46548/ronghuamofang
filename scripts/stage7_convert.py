# -*- coding: utf-8 -*-
"""阶段 7：Markdown → Word（模板底版 + 降级策略，v2 4.6 / P0.6 模板接入）。

- 合并 chapters/refined/*.md → merged/book.md（缺 refined 时回退 checked/raw）；
- 支持用户 Word 模板（templates/*.dotx/.docx，config/project.yaml book.word_template）：
  * .dotx 自动归一为可解析 .docx（content type 转换）；
  * 替换模板占位符：[书籍标题]→书名、[作者]→作者、[目录占位符]→TOC 域（Word 打开 F9 更新）；
  * 正文从 [请输入文本] 锚点之后开始写入；
- 无模板时回退 P0 硬编码方案（宋体正文 + 黑体标题）；
- 复杂元素（表格/图片/代码块）降级为普通段落文本；
- 章节标题按 chapter.page_break 决定是否新起一页（默认 true）。
"""
import argparse
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

import sys
import os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from utils.file_io import read_text, write_text
from utils.config_io import load_config_yaml
from utils.text_norm import normalize_chinese_punctuation_with_counts

HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
LIST_ITEM_RE = re.compile(r"^\s*[-*]\s+(.*)$")
NUM_ITEM_RE = re.compile(r"^\s*\d+[.、]\s+(.*)$")
TABLE_ROW_RE = re.compile(r"^\s*\|.*\|\s*$")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\([^)]*\)")
CODE_FENCE_RE = re.compile(r"^```")
COMMENT_RE = re.compile(r"<!--.*?-->", re.S)

# 模板占位符
PH_TITLE = "[书籍标题]"
PH_AUTHOR = "[作者]"
PH_TOC = "[目录占位符]"
PH_BODY = "[请输入文本]"

TEMPLATE_CACHE = Path("data/tmp/template_cache")


def merge_book(refined_dir, checked_dir, raw_dir, out_path, book_title="未命名"):
    """按优先级合并章节为 book.md：refined > checked > raw，**逐章**取最优版本。

    ⚠️ 2026-10-01 修正：原先「命中第一个非空目录就定死」。于是
    `data/chapters/refined/` 里只要残留**一个**文件（例如上一轮 stage6 只完成了一部分），
    就只合并那一章，checked/raw 里其它章**全被忽略** —— 而成品就此定稿；
    更糟的是 stage7 会 `mark_stage_done`，断点续跑再也不会补它。
    现在逐章按优先级取（refined 优先，缺的降级到 checked/raw），
    合并章数由调用方校验（见 run_stage）。
    """
    files = {}
    for d in (refined_dir, checked_dir, raw_dir):
        p = Path(d)
        if not p.exists():
            continue
        for f in sorted(p.glob("*.md")):
            files.setdefault(f.name, f)          # 先到先得 = 优先级高的赢
    if not files:
        raise FileNotFoundError("无可用章节（refined/checked/raw 均为空）")
    parts = [f"# {book_title}", ""]
    # 剥协议标记：老产物首行可能残留 `===FILE: <绝对路径>===`（见 strip_protocol_markers），
    # 不剥就会直接进 Word 成品（既是脏数据，也把本机绝对路径带出去）。
    from utils.llm_client import strip_protocol_markers
    for name in sorted(files):
        parts.append(strip_protocol_markers(read_text(files[name])).strip())
        parts.append("")
    text = COMMENT_RE.sub("", "\n".join(parts))  # 去除质量/摘要注释
    # 中文标点归一化兜底（2026-10-18）：前面各阶段的归一化只覆盖**新产出**的章，
    # 历史遗留的脏章（修复前写下的）在合并这一步统一清洗 —— 成品 docx 永不带
    # 半角引号 / 直角引号 / 单破折号。放在去注释之后：注释已剥离，不存在保护区误伤。
    text, _counts = normalize_chinese_punctuation_with_counts(text)
    if _counts:
        detail = "、".join(f"{k}={v}" for k, v in _counts.items() if v)
        print(f"[stage7] 合并时标点归一化（{detail}）")
    write_text(out_path, text.strip() + "\n")
    print(f"[stage7] 合并 {len(files)} 章 → {out_path}")
    return len(files)


# ---------- 模板工具 ----------

def dotx_to_docx(src, dst):
    """把 .dotx 归一为 .docx（替换 content type），供 python-docx 打开。"""
    with zipfile.ZipFile(src) as zin:
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(
                        b"wordprocessingml.template.main+xml",
                        b"wordprocessingml.document.main+xml")
                zout.writestr(item, data)
    return dst


def prepare_template(template_path):
    """返回可被 python-docx 打开的模板 docx 路径。.dotx 需先转换（缓存到 data/tmp）。"""
    src = Path(template_path)
    if src.suffix.lower() == ".docx":
        return src
    TEMPLATE_CACHE.mkdir(parents=True, exist_ok=True)
    dst = TEMPLATE_CACHE / (src.stem + ".docx")
    dotx_to_docx(src, dst)
    return dst


def _find_paragraph(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    return None


def _set_paragraph_text(p, text):
    """清空段落并写入新文本（保留段落样式）。"""
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)
    return p


def _insert_toc_field(p):
    """把段落内容替换为 TOC 域（Word 打开后 F9 更新）。"""
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "（打开文档后按 F9 更新目录）"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr, fld_sep, t, fld_end):
        run._element.append(el)


def _set_font(run, name, size, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def _style_font_map(doc):
    """styleId → eastAsia 字体名（沿 basedOn 链向上取第一个有 eastAsia 的）。

    ## 为什么不能硬编码宋体（2026-10-18 用户点破）

    用户原话：「正文 run 显式写确定不会炸吗？模板里可不是全是宋体。」
    实测模板 `小说格式模板V1.3.dotx` 的样式表：正文 Normal=宋体、章标题
    heading1-4=黑体、卷首诗/引用文本=楷体、书名=黑体 —— **三套字体并存**。
    硬编码宋体会把章标题与引用段一起刷成宋体，等于用一个 bug 换另一个 bug。

    所以此处只做一件事：把模板**自己**声明的字体读出来，供逐 run 沿用。
    模板里没有 eastAsia 的样式（如 Quote）返回 None → 调用方不动该 run，
    维持 Word 的原生继承，不会比修复前更糟。
    """
    raw, based = {}, {}
    for st in doc.styles:
        el = st.element
        sid = el.get(qn("w:styleId"))
        if not sid:
            continue
        f = None
        rpr = el.find(qn("w:rPr"))
        if rpr is not None:
            rf = rpr.find(qn("w:rFonts"))
            if rf is not None:
                f = rf.get(qn("w:eastAsia"))
        raw[sid] = f
        b = el.find(qn("w:basedOn"))
        based[sid] = b.get(qn("w:val")) if b is not None else None

    memo = {}

    def resolve(sid, depth=0):
        if not sid or depth > 8 or sid not in raw:
            return None
        if sid in memo:
            return memo[sid]
        f = raw.get(sid) or resolve(based.get(sid), depth + 1)
        memo[sid] = f
        return f

    return {sid: resolve(sid) for sid in raw}


def _apply_cjk_font(run, font):
    """把 run 的 ascii / hAnsi / eastAsia 三处都设为同一个中文字体，并加 hint。

    ## 三处都要设的原因（这是「破折号不是宋体」的直接根因）

    Word 对 U+2014（—）、U+201C/U+201D（“”）、U+2026（…）这类**通用标点**
    按 ascii（西文）字体渲染。旧实现只靠样式继承，样式的 eastAsia=宋体 而
    ascii=Times New Roman —— 于是汉字是宋体、破折号与引号是 Times New Roman，
    视觉上就是用户看到的「破折号字体不对」。
    `hint=eastAsia` 再兜一层：让 Word 对歧义字符优先按中文字体规则处理。
    """
    if not font:
        return
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), font)
    rf.set(qn("w:hAnsi"), font)
    rf.set(qn("w:eastAsia"), font)
    rf.set(qn("w:hint"), "eastAsia")


def _style_font(doc, style_name, font_map, fallback=None):
    """按样式名取其 eastAsia 字体；取不到用 fallback（通常是 Normal 的宋体）。"""
    try:
        sid = doc.styles[style_name].style_id
    except KeyError:
        return fallback
    return font_map.get(sid) or fallback


def _add_paragraph_before(anchor, text, style=None, font=None):
    """在 anchor 段前插入段落。style 不存在时回退 Normal。

    font 给定时把该 run 的三处字体显式写死（见 `_apply_cjk_font`），
    不给则完全维持模板继承（与修复前一致，零回归风险）。
    """
    p = anchor.insert_paragraph_before()
    if style:
        try:
            p.style = style
        except KeyError:
            pass
    run = p.add_run(text)
    if font:
        _apply_cjk_font(run, font)
    return p


def md_to_docx(md_text, out_path, book_title="未命名", author="",
               template_path=None, page_break=True):
    """行解析转换。template_path 提供时以模板为底版，否则 P0 硬编码样式。"""
    use_template = template_path is not None and Path(template_path).exists()

    if use_template:
        doc = Document(str(prepare_template(template_path)))
    else:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12)
        doc.add_heading(book_title, level=0)

    # 字体解析：模板自己声明什么字体就用什么（宋体正文 / 黑体标题 / 楷体引用）
    font_map = _style_font_map(doc)
    try:
        body_font = font_map.get(doc.styles["Normal"].style_id) or "宋体"
    except KeyError:
        body_font = "宋体"

    # 模板占位符填充
    anchor = None
    if use_template:
        p = _find_paragraph(doc, PH_TITLE)
        if p:
            _set_paragraph_text(p, book_title)
        p = _find_paragraph(doc, PH_AUTHOR)
        if p:
            _set_paragraph_text(p, author)
        p = _find_paragraph(doc, PH_TOC)
        if p:
            _insert_toc_field(p)
        anchor = _find_paragraph(doc, PH_BODY)
        if anchor is None:
            anchor = doc.paragraphs[-1]
        _set_paragraph_text(anchor, "")  # 锚点清空，正文插入其前

    in_code = False
    list_items = []
    skip_first_h1 = use_template  # 模板已有书名页，跳过正文开头的 # 书名

    def flush_list():
        nonlocal list_items
        if not list_items:
            return
        for text, bullet in list_items:
            if use_template:
                style = "List Bullet" if bullet else "List Number"
                p = _add_paragraph_before(anchor, "", style=style, font=body_font)
                p.runs[0].text = text if p.runs else text
                if not p.runs:
                    p.add_run(text)
            else:
                p = doc.add_paragraph(style="List Bullet" if bullet else "List Number")
                p.add_run(text)
        list_items = []

    for raw_line in md_text.splitlines():
        line = raw_line.rstrip()
        if COMMENT_RE.match(line):
            continue
        # 代码块围栏：整块降级为段落
        if CODE_FENCE_RE.match(line):
            in_code = not in_code
            if not in_code:
                continue
            flush_list()
            if use_template:
                _add_paragraph_before(anchor, line[3:].strip() or "（代码块已降级为文本）")
            else:
                doc.add_paragraph(line[3:].strip() or "（代码块已降级为文本）")
            continue
        if in_code:
            flush_list()
            if use_template:
                _add_paragraph_before(anchor, line)
            else:
                doc.add_paragraph(line)
            continue
        # 表格行降级
        if TABLE_ROW_RE.match(line):
            flush_list()
            text = "[表格已降级为文本] " + line.strip()
            if use_template:
                _add_paragraph_before(anchor, text)
            else:
                doc.add_paragraph(text)
            continue
        # 图片降级
        if IMAGE_RE.search(line):
            flush_list()
            alt = IMAGE_RE.search(line).group(1)
            text = f"[图片已降级为文本] {alt}"
            if use_template:
                _add_paragraph_before(anchor, text)
            else:
                doc.add_paragraph(text)
            continue
        # 标题
        hm = HEADING_RE.match(line)
        if hm:
            flush_list()
            hashes = len(hm.group(1))
            if skip_first_h1 and hashes == 1:
                skip_first_h1 = False
                continue  # 模板已有书名页
            if use_template:
                # 正文中 ## 为章标题（Heading 1），### 为节标题（Heading 2）
                level = hashes - 1
                style_name = f"Heading {max(1, min(level, 3))}"
                # 章标题沿用该标题样式自己的 eastAsia 字体（模板里是黑体）
                p = _add_paragraph_before(anchor, hm.group(2).strip(), style=style_name,
                                          font=_style_font(doc, style_name, font_map, body_font))
                if hashes == 2 and page_break:
                    p.paragraph_format.page_break_before = True
            else:
                level = min(hashes, 3)
                p = doc.add_heading(level=level)
                run = p.add_run(hm.group(2).strip())
                _set_font(run, "黑体", 16 - level * 2, bold=True, color=(0x1F, 0x3B, 0x57))
            continue
        # 列表
        lm = LIST_ITEM_RE.match(line)
        nm = NUM_ITEM_RE.match(line)
        if lm or nm:
            list_items.append((lm.group(1) if lm else nm.group(1), bool(lm)))
            continue
        # 空行：刷新列表
        if not line.strip():
            flush_list()
            continue
        # 普通段落
        flush_list()
        if use_template:
            # 模板模式：**只写文本，段落格式全继承模板 Normal**（1.15 行距/首行缩进/两端对齐）。
            # 旧实现硬编码 space_after=6pt + line_spacing=1.5 直写每个段落 → 覆盖模板规格，
            # 违反件一验收线「样式继承自模板非硬编码」（2026-10-11 template_check 抓出）。
            # 字体同理：段落格式继承模板，但 run 的三处 rFonts 显式写模板声明的 eastAsia
            # 字体 —— 否则破折号/引号走样式里的西文字体（Times New Roman / Helvetica）。
            _add_paragraph_before(anchor, line.strip(), font=body_font)
        else:
            p = doc.add_paragraph(line.strip())
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            for r in p.runs:
                _apply_cjk_font(r, body_font)
    flush_list()

    doc.save(str(out_path))
    print(f"[stage7] Word 输出: {out_path}"
          + ("（模板底版）" if use_template else "（无模板，硬编码样式）"))


def run_stage(cfg, proj, progress, db, cost, client=None, task_dir=None, run_id=None):
    book = proj.get("book", {})
    book_name = book.get("name", "未命名")
    author = book.get("author", "") or ""
    tmpl = (book.get("word_template") or "").strip()
    template_path = Path(tmpl) if tmpl else None
    page_break = bool((cfg or {}).get("chapter", {}).get("page_break", True))

    out_dir = Path("output")
    out_dir.mkdir(parents=True, exist_ok=True)
    book_md = Path("data/merged/book.md")
    book_docx = out_dir / f"{book_name}_完整版.docx"

    n = merge_book(Path("data/chapters/refined"), Path("data/chapters/checked"),
                   Path("data/chapters/raw"), book_md, book_name)

    # ---- 导出前校验（第一波防御 · 任务四）：不过关不出文件 ----
    from utils.export_check import check_merged_md, check_template, check_docx_output
    expect = int((proj or {}).get("book", {}).get("chapters", 0) or 0) or None
    md_text = read_text(book_md)
    md_errs, md_warns = check_merged_md(md_text, expect_chapters=expect)
    tmpl_errs, tmpl_warns = ([], [])
    if template_path is not None:
        tmpl_errs, tmpl_warns = check_template(template_path)
    for w in md_warns + tmpl_warns:
        print(f"[stage7] 导出警告：{w}")
    all_errs = md_errs + tmpl_errs
    if all_errs:
        progress.set_stage(7, "failed", error="；".join(all_errs[:3]))
        print("[stage7] 导出前校验未通过 —— 不出文件：")
        for e in all_errs:
            print(f"  · {e}")
        # 失败消息必须**点名**具体问题（只报"N 项未通过"等于让人去翻 stdout
        # 才知道拦了什么 —— orchestrator 日志/GUI 只拿到这条消息）。
        _named = "；".join(all_errs[:3]) + ("…" if len(all_errs) > 3 else "")
        return False, (f"stage7 导出前校验未通过（{len(all_errs)} 项）—— "
                       f"已按「校验不过不出文件」拦下：{_named}")

    md_to_docx(md_text, book_docx, book_title=book_name, author=author,
               template_path=template_path, page_break=page_break)

    # ---- 导出后校验（成品层）：不过就删文件 ----
    out_errs, out_warns = check_docx_output(book_docx, md_text=md_text)
    for w in out_warns:
        print(f"[stage7] 成品警告：{w}")
    if out_errs:
        try:
            book_docx.unlink()
        except Exception:                                    # noqa: BLE001
            pass
        progress.set_stage(7, "failed", error="；".join(out_errs[:3]))
        print("[stage7] 成品校验未通过 —— 文件已删除：")
        for e in out_errs:
            print(f"  · {e}")
        _named = "；".join(out_errs[:3]) + ("…" if len(out_errs) > 3 else "")
        return False, (f"stage7 成品校验未通过（{len(out_errs)} 项）—— "
                       f"残次成品已删除：{_named}")

    if not book_docx.exists() or book_docx.stat().st_size == 0:
        progress.set_stage(7, "failed", error="docx 未生成或为空")
        return False, "stage7 docx 生成失败"

    # 章数校验（2026-10-01）：只判 docx 非空不够 —— **缺章的书也是"非空 docx"**。
    # 而 stage7 一旦 mark_done，断点续跑就再也不会补它：状态显示完成，成品是残卷。
    if expect and n < expect:
        progress.set_stage(7, "failed", error=f"合并 {n} 章 < 计划 {expect} 章")
        return False, (f"stage7 失败：只合并到 {n} 章，计划 {expect} 章 —— "
                       "部分章节可能尚未生成，或不在 refined/checked/raw 里")

    progress.mark_stage_done(7)
    print(f"[stage7] 完成：{n} 章 → {book_docx}")
    return True, f"stage7 完成（{n} 章）"


def main():
    parser = argparse.ArgumentParser(description="NovelForge 阶段7：Markdown→Word")
    args = parser.parse_args()
    cfg = load_config_yaml("config/system.yaml")
    proj = load_config_yaml("config/project.yaml")
    from utils.progress_manager import ProgressManager
    progress = ProgressManager("data/state/progress.json")
    ok, msg = run_stage(cfg, proj, progress, None, None)
    print(msg)
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
